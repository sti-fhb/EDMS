"""操作手冊 Word 產生器（TBMS）

以院方交付之「需求規格確認書」為範本開檔，清空內文但**保留其樣式、編號定義、
分節與頁碼設定**，再填入手冊內容 —— 字體、段落、頁碼一律繼承範本，不另調近似值。

產出結構（對齊需求規格確認書）：
    封面 → 文件制/修訂履歷 → 目錄        （第 1 節，不編頁碼）
    壹、… 貳、… 參、…                    （第 2 節，頁碼自 1 起算）

Markdown 標題層級對應：
    #     文件標題（僅供封面，不進內文）
    ##    章 → 壹、貳、參…（Heading 1）
    ###   節 → 一、二、三…（Heading 2）
    ####  小節 → （一）（二）…（Heading 3）
    粗體單行 → 操作子標題

用法：
    python gen_manual.py <手冊.md 或目錄> [-o 輸出目錄]
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
from manual_config import (   # noqa: E402  專案專屬設定一律自本檔讀取
    ORGANISATION, PROJECT_NAME, MANUAL_DOC_NAME as DOC_NAME, AUTHOR,
    DOC_VERSION, MODULE_NAMES, MODULE_DIRS, MANUAL_NAME_RE,
    NON_MANUAL_FILES, NON_MANUAL_DIRS)

# 樣式範本：自院方「需求規格確認書」清空內容而來，只留樣式、分節與頁碼設定。
# ⛔ 不以產出之手冊充當範本：產手冊時需同時讀寫同一檔、Word 開著即失敗，且樣式會逐代退化。
TEMPLATE = Path("docs/manuals/tools/樣式範本.docx")

FIRST_LINE_INDENT = Cm(0.85)   # 範本敘述段之首行縮排
MAX_IMAGE_WIDTH = Cm(15.5)
DOC_INFO_WIDTHS = (Cm(3.81), Cm(6.35))                    # 封面之文件資訊表
HISTORY_WIDTHS = (Cm(2.03), Cm(3.05), Cm(7.62), Cm(3.81))  # 修訂履歷表
FIRST_RELEASE_NOTE = "初次發行"
NOTE_INDENT = Cm(0.5)          # 提醒框：左縮排（小於條列，避免看似條列子項）
MIN_COL_WIDTH = Cm(1.6)        # 欄寬下限：再窄則標題字會被迫逐字斷行
HALF_CHAR_WIDTH = Pt(6.5)      # 半形字寬（中文字為其兩倍），用於估算欄所需寬度
CELL_PADDING = Cm(0.4)         # 儲存格左右內距合計
CAPTION_SIZE = Pt(10)          # 圖說：小於內文、置中
CAPTION_COLOR = RGBColor(0x40, 0x40, 0x40)

H1_NUMERALS = ["壹", "貳", "參", "肆", "伍", "陸", "柒", "捌", "玖", "拾"]
H2_NUMERALS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
               "十一", "十二", "十三", "十四", "十五", "十六", "十七", "十八"]

HEADING_RE = re.compile(r"^(?P<level>#{1,4})\s+(?P<text>.+)$")
IMAGE_RE = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<src>.+?)\)$")
LIST_RE = re.compile(r"^(?P<indent>\s*)(?P<bullet>[-*]|\d+\.)\s+(?P<text>.+)$")
TABLE_RE = re.compile(r"^\s*\|(.+)\|\s*$")
BOLD_ONLY_RE = re.compile(r"^\*\*(?P<text>[^*]+)\*\*$")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


COVER_ORG_SIZE = 36        # 院名
COVER_DOC_SIZE = 28        # 購案名、文件名、模組名
COVER_OPERATION_SIZE = 22  # 作業名（次於模組名）


# ---- 範本處理 ----------------------------------------------------------

def open_template() -> tuple[Document, object]:
    """開啟範本並清空內文，回傳 (doc, 第1節結束之段落元素)。

    第 1 節的 sectPr 寄生在某個段落的 pPr 內，刪掉該段落會連帶失去分節與
    頁碼設定，故僅清空其內容、保留元素本身作為兩節之分界。
    """
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    section_break_p = None

    for child in list(body):
        if child.tag == qn("w:sectPr"):      # 第 2 節設定，保留
            continue
        if child.tag == qn("w:p"):
            pPr = child.find(qn("w:pPr"))
            if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
                section_break_p = child
                for node in list(child):     # 只留 pPr，清掉文字
                    if node.tag != qn("w:pPr"):
                        child.remove(node)
                continue
        body.remove(child)

    if section_break_p is None:
        raise RuntimeError("範本中找不到分節設定，無法沿用頁碼配置")

    _enable_field_update(doc)
    _ensure_bullet_numbering(doc)
    return doc, section_break_p


def _enable_field_update(doc: Document) -> None:
    """讓 Word 開檔時自動更新目錄頁碼，免去手動按 F9。"""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def _move_before(element, paragraph):
    """把新建的段落／表格搬到第 1 節（分節符之前）。"""
    element.addprevious(paragraph._element if hasattr(paragraph, "_element") else paragraph._p)


# 分頁一律設段落之 page_break_before。
# ⛔ 不用「空段落＋分頁符」：該段落自身佔一行，前頁剛好排滿時會多出一整頁空白。


# ---- 內容元件 ----------------------------------------------------------

def _add_runs(paragraph, text: str) -> None:
    """處理 **粗體** 與 `行內程式碼`。字型一律沿用樣式，不逐一指定。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


BULLET_ABSTRACT_ID = 900       # 另起號段，避免與範本既有之定義相撞
BULLET_NUM_ID = 900


def _ensure_bullet_numbering(doc: Document) -> None:
    """為無序清單補一組項目符號之編號定義。

    ⚠️ 範本僅有 decimal（「1.」）一種定義，List Bullet 樣式亦未關聯任何編號——
    不補此定義，無序清單即無符號可用，只能以文字替代。
    """
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    numbering = doc.part.part_related_by(RT.NUMBERING).element
    if numbering.find(qn("w:num") + f"[@{qn('w:numId')}='{BULLET_NUM_ID}']") is not None:
        return

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(BULLET_ABSTRACT_ID))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"),
                     ("w:lvlText", "●"), ("w:lvlJc", "left")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    pPr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "482")       # 0.85cm
    ind.set(qn("w:hanging"), "482")
    pPr.append(ind)
    lvl.append(pPr)
    abstract.append(lvl)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(BULLET_NUM_ID))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(BULLET_ABSTRACT_ID))
    num.append(ref)

    # abstractNum 一律置於全部 num 之前（schema 要求）
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    numbering.append(num)


def _apply_numbering(paragraph, num_id: int) -> None:
    """掛上編號定義。

    ⚠️ 範本之 List Number 樣式僅承載縮排，**未關聯任何編號定義**——只套樣式而不掛
    numPr，產出將完全不顯示數字。
    """
    pPr = paragraph._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numPr.append(ilvl)
    numPr.append(num)
    pPr.append(numPr)


def _add_list_item(doc: Document, bullet: str, text: str, depth: int, num_id: int | None):
    """條列採範本既有之清單樣式；有序者另掛編號定義。

    縮排與凸排由樣式承載，⛔ 不手動設定——版面日後要調，改樣式即可套用全篇。
    有序清單之 num_id 由呼叫端按「組」配發：同一組共用一個編號定義，故中間夾提醒框
    或圖片時編號仍延續；換組則另配一個，不與前組累加。
    """
    ordered = bullet[0].isdigit()
    base = "List Number" if ordered else "List Bullet"
    style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
    para = doc.add_paragraph(style=style)
    _apply_numbering(para, num_id if ordered and num_id else BULLET_NUM_ID)
    _add_runs(para, text)
    return para


def _style_note(paragraph) -> None:
    """提醒框樣式：左縮排＋上下留白。

    ⚠️ 左縮排 MUST 小於清單樣式之縮排，否則提醒框將被誤認為條列之子項。
    """
    pf = paragraph.paragraph_format
    pf.left_indent = NOTE_INDENT
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)


def _add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    for tag, attr, text in (
        ("w:fldChar", ("w:fldCharType", "begin"), None),
        ("w:instrText", ("xml:space", "preserve"), 'TOC \\o "1-2" \\h \\z \\u'),
        ("w:fldChar", ("w:fldCharType", "separate"), None),
        ("w:t", None, "目錄將於開啟文件時自動更新"),
        ("w:fldChar", ("w:fldCharType", "end"), None),
    ):
        el = OxmlElement(tag)
        if attr:
            el.set(qn(attr[0]), attr[1])
        if text:
            el.text = text
        run._r.append(el)


def build_front_matter(doc: Document, break_p, title: str, module: str | None,
                       doc_name: str = DOC_NAME) -> None:
    """封面、修訂履歷、目錄——一律插入第 1 節。

    測試報告與操作手冊共用本函式，僅文件名稱（doc_name）不同。
    """
    def cover_line(text: str = "", size: int | None = None):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        if text:
            run = para.add_run(text)
            run.font.size = Pt(size)
        _move_before(break_p, para)

    cover_line()
    cover_line(ORGANISATION, COVER_ORG_SIZE)
    for _ in range(8):
        cover_line()
    cover_line(PROJECT_NAME, COVER_DOC_SIZE)
    cover_line(doc_name, COVER_DOC_SIZE)
    for _ in range(7):
        cover_line()
    if module:                                   # 模組名為主、作業名次之
        cover_line(module, COVER_DOC_SIZE)
        cover_line(title, COVER_OPERATION_SIZE)
    else:                                        # 不在模組資料夾下者，封面只印文件名
        cover_line(title, COVER_DOC_SIZE)
    for _ in range(5):                           # 與文件資訊表之間距
        cover_line()

    today = date.today().strftime("%Y/%m/%d")

    # 封面頁之文件資訊（範本表 0：二欄、置中）
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in (("版本", DOC_VERSION), ("建立日期", today),
                         ("更新日期", ""), ("作者", AUTHOR)):
        cells = info.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell, width in zip(cells, DOC_INFO_WIDTHS):
            cell.width = width
    _move_before(break_p, info)      # 文件資訊表接於封面下方，同頁不分頁

    heading = doc.add_paragraph("文件制/修訂履歷", style="Heading 1 No TOC")
    heading.paragraph_format.page_break_before = True
    _move_before(break_p, heading)

    # 修訂履歷（範本表 1：四欄橫式，表頭粗體置中；說明欄資料靠左）
    history = doc.add_table(rows=0, cols=4)
    history.style = "Table Grid"
    header = history.add_row().cells
    for cell, text, width in zip(header, ("版次", "日期", "說明", "修訂人"), HISTORY_WIDTHS):
        cell.width = width
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run(text).bold = True

    row = history.add_row().cells
    for i, (cell, text, width) in enumerate(
            zip(row, (DOC_VERSION, today, FIRST_RELEASE_NOTE, AUTHOR), HISTORY_WIDTHS)):
        cell.width = width
        cell.text = text
        if i != 2:                                   # 說明欄靠左，其餘置中
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _move_before(break_p, history)

    heading = doc.add_paragraph("目錄", style="Heading 1 No TOC")
    heading.paragraph_format.page_break_before = True
    _move_before(break_p, heading)
    toc = doc.add_paragraph()
    _add_toc_field(toc)
    _move_before(break_p, toc)


# ---- Markdown 轉換 ------------------------------------------------------

def _text_width(text: str) -> float:
    """估算文字排出來有多寬：中文字約佔兩個英數字寬。"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in re.sub(r"\*\*|`|<br>", "", text))


def _column_widths(parsed: list[list[str]], total: int) -> list[int]:
    """依各欄內容長短分配欄寬，總和填滿版面寬度。

    ⚠️ 不分配即為 Word 預設之等寬——「欄位」這類短標題欄會與整段說明同寬，
    左半留下大片空白。

    兩道規則：
    1. 按內容長度之平方根分配。取平方根是因長說明欄本來就會換行，不需按最長
       句給足寬度；不壓縮則兩欄之比可達 7:1，短欄會被擠到容不下標題。
    2. **一欄不必比它最長的內容還寬**——放得下就固定在該寬度，省下的部分讓給
       仍需換行的欄。只做第 1 條的話，「查詢範圍」這種四字欄仍會分到 4.7cm。
    """
    cols = len(parsed[0])
    longest = [max((_text_width(row[i]) for row in parsed if i < len(row)), default=1)
               for i in range(cols)]
    # 內容排完所需之寬度（含儲存格左右內距）
    ideal = [max(int(w * HALF_CHAR_WIDTH) + CELL_PADDING, MIN_COL_WIDTH) for w in longest]
    weights = [max(w, 1) ** 0.5 for w in longest]

    widths = [0] * cols
    settled: set[int] = set()
    while True:
        spare = total - sum(widths[i] for i in settled)
        pool = sum(weights[i] for i in range(cols) if i not in settled)
        newly = {i for i in range(cols) if i not in settled
                 and ideal[i] <= spare * weights[i] / pool}
        for i in range(cols):
            if i not in settled:
                widths[i] = ideal[i] if i in newly else int(spare * weights[i] / pool)
        if not newly or len(settled | newly) == cols:
            break
        settled |= newly

    # 各欄內容都放得下時總寬會短於版面，餘寬按權重分回，使每張表左右對齊
    if (short := total - sum(widths)) > 0:
        widths = [w + int(short * k / sum(weights)) for w, k in zip(widths, weights)]
    widths[-1] += total - sum(widths)      # 吸收整數捨入之差額
    return widths


def _set_column_widths(table, widths: list[int]) -> None:
    """套用欄寬。

    ⚠️ Word 只認 fixed 版面配置下的欄寬；留在 autofit 會依內容重算而蓋掉設定。
    且欄寬記在每個儲存格上，MUST 逐格設定，只設首列無效。
    """
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def _add_table(doc: Document, rows: list[str]) -> None:
    parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = parsed[0], parsed[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"

    for cell, text in zip(table.rows[0].cells, header):
        cell.paragraphs[0].text = ""
        _add_runs(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.paragraphs[0].text = ""
            for i, line in enumerate(text.split("<br>")):
                para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                _add_runs(para, line)

    section = doc.sections[-1]
    usable = section.page_width - section.left_margin - section.right_margin
    _set_column_widths(table, _column_widths([header] + body, usable))


def convert_body(doc: Document, md_path: Path) -> str:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    title = md_path.stem
    counters = {1: 0, 2: 0, 3: 0}
    list_num_id = 0          # 清單組序；範本備有多組編號定義可供配發

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line or (set(line) == {"-"} and line.startswith("---")):
            i += 1
            continue

        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            if len(block) >= 2:
                _add_table(doc, block)
            continue

        if line.lstrip().startswith(">"):   # 含縮排於條列內之提醒框
            block = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                text = lines[i].lstrip()[1:].strip()
                if text:
                    block.append(text)
                i += 1
            for text in block:
                if item := re.match(r"^[-*]\s+(?P<text>.+)$", text):
                    # 提醒框內之條列：掛項目符號，縮排與提醒框對齊
                    para = doc.add_paragraph()
                    _apply_numbering(para, BULLET_NUM_ID)
                    _add_runs(para, item.group("text"))
                    para.paragraph_format.left_indent = NOTE_INDENT + Cm(0.6)
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                else:
                    para = doc.add_paragraph()
                    _add_runs(para, text)
                    _style_note(para)
            continue

        if match := HEADING_RE.match(line):
            level = len(match.group("level"))
            text = match.group("text").strip()
            if level == 1:                      # 文件標題只供封面
                title = text
                i += 1
                continue
            depth = level - 1                   # ## → 1, ### → 2, #### → 3
            counters[depth] += 1
            for deeper in range(depth + 1, 4):
                counters[deeper] = 0
            if depth == 1:
                if counters[1] > len(H1_NUMERALS):
                    raise ValueError(
                        f"{md_path.name} 之章數超過 {len(H1_NUMERALS)}；"
                        f"手冊固定三章，請檢查是否誤將非本流程之文件納入")
                label = f"{H1_NUMERALS[counters[1] - 1]}、{text}"
            elif depth == 2:
                label = f"{H2_NUMERALS[counters[2] - 1]}、{text}"
            else:
                label = f"（{H2_NUMERALS[counters[3] - 1]}）{text}"
            doc.add_paragraph(label, style=f"Heading {depth}")
            i += 1
            continue

        if match := IMAGE_RE.match(line):
            src = (md_path.parent / match.group("src")).resolve()
            if src.exists():
                doc.add_picture(str(src), width=MAX_IMAGE_WIDTH)
                picture = doc.paragraphs[-1]
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.space_before = Pt(8)
                picture.paragraph_format.space_after = Pt(2)
                picture.paragraph_format.left_indent = Cm(0)
                picture.paragraph_format.first_line_indent = Cm(0)
                if caption := match.group("alt"):
                    para = doc.add_paragraph()
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_after = Pt(8)
                    para.paragraph_format.first_line_indent = Cm(0)
                    run = para.add_run(caption)
                    run.font.size = CAPTION_SIZE
                    run.font.color.rgb = CAPTION_COLOR
            else:
                doc.add_paragraph(f"［找不到圖片：{src.name}］")
            i += 1
            continue

        if match := BOLD_ONLY_RE.match(line):   # 整行粗體視為操作子標題
            doc.add_paragraph(match.group("text"), style="操作子標題")
            i += 1
            continue

        if match := LIST_RE.match(line):
            bullet = match.group("bullet")
            if bullet[0].isdigit() and bullet.startswith("1."):
                list_num_id += 1        # 自 1. 起算者為新的一組
            _add_list_item(doc, bullet, match.group("text"),
                           depth=len(match.group("indent")) // 2,
                           num_id=list_num_id if bullet[0].isdigit() else None)
            i += 1
            continue

        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = FIRST_LINE_INDENT
        _add_runs(para, line)
        i += 1

    return title


def _extract_module(md_path: Path) -> str | None:
    """依所在資料夾判定所屬模組；不在模組資料夾下者回 None。"""
    return MODULE_NAMES.get(md_path.parent.name)


def _verify_output(md_path: Path, out_path: Path) -> list[str]:
    """比對產出之 Word 是否涵蓋 md 的每一項可見文字。

    產生器對無法辨識的語法一律略過而不報錯（圖說曾因此整批遺失且無人察覺），
    故產出後自我核對一次，缺漏即回報。
    """
    doc = Document(str(out_path))
    haystack = "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )

    missing = []
    for line in md_path.read_text(encoding="utf-8").splitlines():
        text = line.strip()
        if not text or set(text) <= set("-|: "):
            continue
        if text.startswith("# "):                     # 文件標題僅出現於封面
            continue
        if text.startswith("#"):
            candidates = [re.sub(r"^#+\s*", "", text)]
        elif text.startswith(">"):
            inner = re.sub(r"^>\s*", "", text)
            candidates = [re.sub(r"^[-*]\s+", "", inner)]   # 含提醒框內之條列
        elif text.startswith("!["):
            candidates = [re.match(r"!\[(.*?)\]", text).group(1)]
        elif text.startswith("|"):
            candidates = [c.strip() for c in text.strip("|").split("|")]
        else:
            # 條列符號後必有空白；⛔ 不可寫成 \s*，否則整行粗體之首個 * 會被當條列符號吃掉
            candidates = [re.sub(r"^(\d+\.|[-*])\s+", "", text)]

        for candidate in candidates:
            candidate = re.sub(r"\*\*(.+?)\*\*", r"\1", candidate)
            candidate = re.sub(r"`(.+?)`", r"\1", candidate).strip()
            if candidate and candidate not in haystack:
                missing.append(candidate)
    return missing


def build_manual(md_path: Path, out_path: Path) -> None:
    doc, break_p = open_template()
    title = convert_body(doc, md_path)          # 先轉內文以取得文件標題
    build_front_matter(doc, break_p, title, _extract_module(md_path))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    missing = _verify_output(md_path, out_path)
    if missing:
        print(f"已產出：{out_path}　但有 {len(missing)} 項內容未進入 Word：")
        for item in missing[:5]:
            print(f"    {item[:60]}")
    else:
        print(f"已產出：{out_path}（內容核對通過）")


def main() -> None:
    parser = argparse.ArgumentParser(description="操作手冊 Markdown → Word（套用院方文件格式）")
    parser.add_argument("source", type=Path, help="手冊 .md 檔或目錄")
    parser.add_argument("-o", "--output", type=Path, help="輸出目錄（預設與來源同層）")
    args = parser.parse_args()

    if not TEMPLATE.exists():
        raise SystemExit(f"找不到格式範本：{TEMPLATE}")

    if args.source.is_dir():
        candidates = [md for md in sorted(args.source.rglob("*.md"))
                      if not any(part.startswith("_") or part in NON_MANUAL_DIRS
                                 for part in md.parent.parts)
                      and md.name not in NON_MANUAL_FILES]
        sources, skipped = [], []
        for md in candidates:
            # 模組目錄下之檔案須符合命名慣例，否則非本流程產出
            if md.parent.name in MODULE_DIRS and not MANUAL_NAME_RE.match(md.stem):
                skipped.append(md)
            else:
                sources.append(md)
        for md in skipped:
            print(f"略過：{md.relative_to(args.source)}（檔名不符「作業碼-作業名」慣例）")
    else:
        sources = [args.source]
    for md in sources:
        out_dir = args.output or md.parent
        build_manual(md, out_dir / f"{md.stem}.docx")


if __name__ == "__main__":
    main()
