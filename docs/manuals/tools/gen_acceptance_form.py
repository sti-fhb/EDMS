"""測試報告產生器（TBMS 操作手冊配套）

從 docs/manuals/_tests/ 之測試項目來源檔，產出測試報告 Word 檔。

交付時第一張為**總表**（院方附表 7：各模組一列，測試人員勾選合格／不合格），
其後為各模組之測試報告。

**一個模組一個來源檔、一份 Word**：封面與目錄（共用樣式範本，不印文件資訊表與修訂履歷），
其後各作業獨立一節，只列「項次｜測試情境」，項次於各表內自 1 起算（2026-10-29 BC／TL SA 裁示）。
項目表中「項次空白」之列視為分段標題（如依站點分「一、用血單位」），產出時合併為整列。

用法：
    python gen_acceptance_form.py            # 產出 _tests/ 下全部模組與總表
    python gen_acceptance_form.py <檔案.md>  # 單一模組
    python gen_acceptance_form.py --summary  # 只產總表

MEMO 版（`_tests/_memo/*_MEMO.md`，內部協測備忘、非交付，只供測試當下參考）：封面只留標題與
目錄（不產文件資訊表與修訂履歷），保留測試注意事項；各作業之受測資訊表只留標題列
「測試情境說明」（不留購案／測試單位／測試日期三列），項目表為「項次｜情境｜MEMO」三欄；
預設輸出於來源檔同資料夾，不進交付之測試報告目錄。

⚠️ 測試項目與操作手冊分開維護：手冊供操作人員照做，本表供驗收測試，讀者與用途皆不同。
⚠️ 測試內容唯一來源為 _tests/ 下之 md，⛔ 不在本檔另寫一份——兩份必然漂移。
⚠️ 表格格式依院方提供之附表 7 圖樣重刻（無原始範本檔）；取得範本後應改以範本套版。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

sys.path.insert(0, str(Path(__file__).parent))
from gen_manual import (   # noqa: E402
    CELL_SPACING, FIRST_LINE_INDENT, build_front_matter, finalize_fields, open_template)
from manual_config import (   # noqa: E402  專案專屬設定一律自本檔讀取
    PROJECT_NAME, REPORT_DOC_NAME as DOC_NAME, FORM_TITLE, FORM_CAPTION, MODULE_NAMES,
    REPORT_ENV_NOTICE, REPORT_ENV_NOTICE_TITLE, REPORT_SUMMARY_MODULES, REPORT_SUMMARY_EDMS)

FORM_FONT = "標楷體"

# 兩張表各自的欄寬：基本資料表與項目表欄位意義不同，不共用寬度
INFO_WIDTHS = (Cm(2.6), Cm(6.4), Cm(3.4), Cm(4.6))
ITEM_WIDTHS = (Cm(1.6), Cm(14.9))
MEMO_WIDTHS = (Cm(1.6), Cm(8.4), Cm(7.0))
MEMO_DIR = "_memo"
MEMO_TITLE = "測試情境說明"   # MEMO 版受測資訊表只留此標題列（2026-10-29 user 裁示）
DEFAULT_OUTPUT = Path("docs/manuals/測試報告")
SUMMARY_NAME = "00-測試報告總表.docx"   # 00 開頭：檔案總管中排在各模組之前

# 總表為單一張表，基本資料與項目之欄位界線不同，故以六欄為底、各列自行合併：
# 項次｜購案名稱等標籤之右半｜測試內容｜單位主管用印等標籤｜合格｜不合格
SUMMARY_WIDTHS = (Cm(1.5), Cm(1.0), Cm(6.7), Cm(3.1), Cm(2.3), Cm(1.9))
SUMMARY_ROW_HEIGHT = Cm(1.2)

OPERATION_RE = re.compile(r"^##\s+(?P<name>\S+\s+.+?)\s*$")   # 「## BS01 入庫作業」
ITEM_SECTION = re.compile(r"^###\s+測試項目清單\s*$")
TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
UNDERLINE_RE = re.compile(r"<u>(.+?)</u>")   # 測試資料代稱（如<u>收血單B</u>）印成底線
# 行內程式碼：只去反引號、內文照排（EDMS #802；與 gen_manual 之 INLINE_RE 同一組語法）。
# 先去粗體再去反引號：粗體內可再含行內程式碼（如 **先抄下原值 `0 8 * * *`**）
CODE_RE = re.compile(r"`(.+?)`")


def _spacing_element(twips: int):
    el = OxmlElement("w:spacing")
    el.set(qn("w:val"), str(twips))
    return el


def _style_run(run, *, size: int = 12, bold: bool = False, spacing: bool = False):
    run.font.name = FORM_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FORM_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if spacing:  # 標題字距加寬，貼近院方表頭樣式
        run.font.element.rPr.append(_spacing_element(80))
    return run


def _set_cell(cell, text: str, *, size: int = 12, bold: bool = False,
              align=WD_ALIGN_PARAGRAPH.CENTER, spacing: bool = False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    para.paragraph_format.space_before = CELL_SPACING   # 與手冊表格同一值
    para.paragraph_format.space_after = CELL_SPACING
    pos = 0
    for match in UNDERLINE_RE.finditer(text):
        if match.start() > pos:
            _style_run(para.add_run(text[pos:match.start()]), size=size, bold=bold, spacing=spacing)
        _style_run(para.add_run(match.group(1)), size=size, bold=bold, spacing=spacing).underline = True
        pos = match.end()
    if pos < len(text) or not text:
        _style_run(para.add_run(text[pos:]), size=size, bold=bold, spacing=spacing)
    return cell


def _apply_widths(row, widths) -> None:
    for cell, width in zip(row.cells, widths):
        cell.width = width


def _repeat_as_header(row) -> None:
    """標記為標題列，跨頁時於每頁頂端重複。

    ⚠️ Word 僅重複「表格最前面連續」之標題列，故基本資料與項目清單 MUST 拆為兩張
    表格——否則欲重複項目表頭，連購案名稱與用印欄都會跟著印在每一頁。
    """
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def parse_module(md_path: Path) -> tuple[str, list[tuple[str, list[tuple[str, str, str]]]]]:
    """回傳 (模組名稱, [(作業名稱, [(項次, 測試內容, MEMO), ...]), ...])；無 MEMO 欄者為空字串。"""
    lines = md_path.read_text(encoding="utf-8").splitlines()

    module = md_path.stem
    for line in lines:
        if line.startswith("# "):
            module = line[2:].strip()
            break

    operations: list[tuple[str, list[tuple[str, str, str]]]] = []
    in_items = False

    for line in lines:
        if match := OPERATION_RE.match(line):
            operations.append((match.group("name"), []))
            in_items = False
            continue
        if ITEM_SECTION.match(line):
            in_items = True
            continue
        if line.startswith("#"):            # 其他標題即結束項目區
            in_items = False
            continue
        if not in_items or not operations:
            continue
        if match := TABLE_ROW.match(line):
            cells = [c.strip() for c in match.group(1).split("|")]
            if len(cells) < 2:
                continue
            no, content = cells[0], cells[1]
            memo = cells[2] if len(cells) > 2 else ""
            if no == "項次" or (no and set(no) <= set("-: ")):   # 表頭與分隔列
                continue
            if not no and not content:
                continue
            # 項次空白之列為分段標題（如「一、用血單位」），產出時合併整列
            operations[-1][1].append((no, CODE_RE.sub(lambda m: m.group(1), BOLD_RE.sub(lambda m: m.group(1), content)),
                                      memo.replace("<br>", "\n")))

    operations = [(name, items) for name, items in operations if items]
    if not operations:
        raise ValueError(f"{md_path.name} 無任何作業之「### 測試項目清單」，不產出測試報告")
    return module, operations


def _add_info_table(doc, *, title: str = FORM_TITLE, with_rows: bool = True) -> None:
    """表一：表頭與受測資訊（不跨頁重複）；with_rows=False 只留標題列（MEMO 版）。"""
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    row = table.add_row()
    _apply_widths(row, INFO_WIDTHS)
    _set_cell(row.cells[0].merge(row.cells[3]), title, size=16, bold=True, spacing=True)
    if not with_rows:
        return

    for left_label, left_value, right_label in (
        ("購案名稱", PROJECT_NAME, "單位主管用印"),
        ("測試單位", "", "測試人員"),
        ("測試日期", "", "測試完畢日期"),
    ):
        row = table.add_row()
        _apply_widths(row, INFO_WIDTHS)
        _set_cell(row.cells[0], left_label)
        _set_cell(row.cells[1], left_value)
        _set_cell(row.cells[2], right_label)
        _set_cell(row.cells[3], "")


def _add_memo_table(doc, items: list[tuple[str, str, str]]) -> None:
    """MEMO 版項目表：項次／情境／MEMO 三欄，表頭跨頁重複。"""
    table = doc.add_table(rows=0, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    head = table.add_row()
    _apply_widths(head, MEMO_WIDTHS)
    for cell, text in zip(head.cells, ("項次", "情境", "MEMO")):
        _set_cell(cell, text)
    _repeat_as_header(head)

    for no, content, memo in items:
        row = table.add_row()
        _apply_widths(row, MEMO_WIDTHS)
        if not no:   # 分段標題列
            _set_cell(row.cells[0].merge(row.cells[2]), content, size=11, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT)
            continue
        _set_cell(row.cells[0], no, size=11)
        _set_cell(row.cells[1], content, size=11, align=WD_ALIGN_PARAGRAPH.LEFT)
        _set_cell(row.cells[2], memo, size=11, align=WD_ALIGN_PARAGRAPH.LEFT)


def _add_item_table(doc, items: list[tuple[str, str, str]]) -> None:
    """測試項目：項次｜測試情境兩欄，表頭跨頁重複。合格與否改於總表逐模組勾選。"""
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    head = table.add_row()
    _apply_widths(head, ITEM_WIDTHS)
    _set_cell(head.cells[0], "項次")
    _set_cell(head.cells[1], "測試情境")
    _repeat_as_header(head)

    for no, content, _memo in items:
        row = table.add_row()
        _apply_widths(row, ITEM_WIDTHS)
        if not no:   # 分段標題列
            _set_cell(row.cells[0].merge(row.cells[1]), content, size=11, bold=True,
                      align=WD_ALIGN_PARAGRAPH.LEFT)
            continue
        _set_cell(row.cells[0], no, size=11)
        _set_cell(row.cells[1], content, size=11, align=WD_ALIGN_PARAGRAPH.LEFT)


def _add_env_notice(doc) -> None:
    """開頭之執行環境警語——目錄之後、第一支作業之前。

    ⚠️ 加在本檔而非 `build_front_matter`：後者為**操作手冊與測試報告共用**，
    改那裡會讓操作手冊也長出這段（手冊是給正式環境的操作人員看的）。
    """
    doc.add_paragraph(REPORT_ENV_NOTICE_TITLE, style="Heading 1")
    for line in REPORT_ENV_NOTICE.splitlines():
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = FIRST_LINE_INDENT   # 比照手冊內文首行縮排
        _style_run(para.add_run(line), size=12)   # ⛔ 不加寬字距：該樣式僅供表頭標題


def _add_operation(doc, operation: str, items: list[tuple[str, str, str]], first: bool,
                   memo: bool = False) -> None:
    heading = doc.add_paragraph(operation, style="Heading 1")
    if not first:
        heading.paragraph_format.page_break_before = True
    if memo:
        _add_info_table(doc, title=MEMO_TITLE, with_rows=False)
        doc.add_paragraph()          # 兩表之間留一行，避免框線相黏
        _add_memo_table(doc, items)
    else:
        _add_item_table(doc, items)


def build_report(module: str, operations: list[tuple[str, list[tuple[str, str, str]]]],
                 out_path: Path, source: Path, memo: bool = False) -> None:
    doc, break_p = open_template()
    # front matter 另以 _move_before 插到第 1 節，不受此處之順序影響。
    # MEMO 版之警語須在全部作業之前；⚠️ 此時各作業一律分頁，否則第一支會黏在警語後面
    if memo:
        _add_env_notice(doc)
    for index, (operation, items) in enumerate(operations):
        _add_operation(doc, operation, items, first=index == 0 and not memo, memo=memo)
    # MEMO 版非正式文件（2026-10-29 user 裁示），與正式報告同樣不產文件資訊表與修訂履歷
    build_front_matter(doc, break_p, module, None, doc_name=DOC_NAME, source=source,
                       with_history=False)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    finalize_fields(out_path)
    detail = "、".join(f"{name} {sum(1 for no, *_ in items if no)} 項" for name, items in operations)
    print(f"已產出：{out_path}（{len(operations)} 支作業：{detail}）")


def _summary_items() -> list[str]:
    return [f"{code.upper()}-{MODULE_NAMES[code]}" for code in REPORT_SUMMARY_MODULES] + \
        list(REPORT_SUMMARY_EDMS)


def _summary_row(table, *, fixed_height: bool = True):
    row = table.add_row()
    _apply_widths(row, SUMMARY_WIDTHS)
    if fixed_height:
        row.height = SUMMARY_ROW_HEIGHT
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return row


def build_summary(out_path: Path) -> None:
    """總表：附表 7 之完整表頭，各模組一列供勾選合格／不合格，單頁。"""
    if not _summary_items():   # EDMS 端不產總表（由 TBMS 端產出，見 manual_config.REPORT_SUMMARY_*）
        return
    doc, break_p = open_template()
    # 只有一頁：改用第 1 節（封面，無頁碼）之版面並移除分節，免得多出一頁空白
    body = doc.element.body
    body.replace(body.find(qn("w:sectPr")), break_p.find(qn("w:pPr")).find(qn("w:sectPr")))
    body.remove(break_p)

    caption = doc.add_paragraph()
    _style_run(caption.add_run(FORM_CAPTION), size=12, bold=True)

    table = doc.add_table(rows=0, cols=len(SUMMARY_WIDTHS))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    row = _summary_row(table)
    _set_cell(row.cells[0].merge(row.cells[5]), FORM_TITLE, size=16, bold=True, spacing=True)
    for left_label, left_value, right_label in (
        ("購案名稱", PROJECT_NAME, "單位主管用印"),
        ("測試單位", "", "測試人員"),
        ("測試日期", "", "測試完畢日期"),
    ):
        row = _summary_row(table)
        _set_cell(row.cells[0].merge(row.cells[1]), left_label)
        _set_cell(row.cells[2], left_value)
        _set_cell(row.cells[3], right_label)
        _set_cell(row.cells[4].merge(row.cells[5]), "")

    head_top = _summary_row(table, fixed_height=False)
    head_bottom = _summary_row(table, fixed_height=False)
    _set_cell(head_top.cells[0].merge(head_bottom.cells[0]), "項次")
    _set_cell(head_top.cells[1].merge(head_bottom.cells[3]), "測試內容")
    _set_cell(head_top.cells[4].merge(head_top.cells[5]), "測試人員勾選")
    _set_cell(head_bottom.cells[4], "合格")
    _set_cell(head_bottom.cells[5], "不合格")

    items = _summary_items()
    for no, name in enumerate(items, 1):
        row = _summary_row(table)
        _set_cell(row.cells[0], str(no))
        _set_cell(row.cells[1].merge(row.cells[3]), name, align=WD_ALIGN_PARAGRAPH.LEFT)
        _set_cell(row.cells[4], "")
        _set_cell(row.cells[5], "")

    doc.core_properties.title = f"{DOC_NAME}總表"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"已產出：{out_path}（{len(items)} 個模組）")


def main() -> None:
    parser = argparse.ArgumentParser(description="測試項目 → 測試報告")
    parser.add_argument("source", type=Path, nargs="?", default=Path("docs/manuals/_tests"),
                        help="測試項目 .md 檔或目錄（預設 docs/manuals/_tests）")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT,
                        help="輸出目錄（預設 docs/manuals/測試報告）")
    parser.add_argument("--summary", action="store_true", help="只產出總表")
    args = parser.parse_args()

    if args.summary or args.source.is_dir():
        build_summary(args.output / SUMMARY_NAME)
    if args.summary:
        return

    # MEMO 版為內部協測備忘，整批產出時略過，須指名單檔產出
    sources = ([p for p in sorted(args.source.rglob("*.md")) if MEMO_DIR not in p.parts]
               if args.source.is_dir() else [args.source])
    for md in sources:
        try:
            module, operations = parse_module(md)
        except ValueError as exc:
            print(f"略過：{exc}")
            continue
        is_memo = md.stem.endswith("_MEMO")
        out_dir = md.parent if is_memo and args.output == DEFAULT_OUTPUT else args.output
        build_report(module, operations, out_dir / f"{md.stem}.docx", md, memo=is_memo)


if __name__ == "__main__":
    main()
